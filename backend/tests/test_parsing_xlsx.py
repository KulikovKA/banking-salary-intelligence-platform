"""Risk-focused tests for the bounded offline XLSX parser."""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from app.contracts.parsing import ParseRequest
from app.parsing.xlsx import PARSER_VERSION, XlsxParser


def _request(path: Path, *, parser_version: str = PARSER_VERSION) -> ParseRequest:
    return ParseRequest(
        source_version_id=uuid4(),
        raw_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        kind="xlsx",
        corpus_kind="synthetic",
        parser_version=parser_version,
    )


def _sample_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Тарифы"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Зарплатный проект"
    sheet.append(["Услуга", "Комиссия"])
    sheet.append(["Перевод", 0.003])
    sheet["B3"].number_format = "0.0%"
    sheet["C3"] = "0,3%"
    sheet.append(["Обслуживание", "Индивидуально"])
    sheet["A5"] = "Условие [1]"
    sheet["B5"] = "=1+1"
    sheet["A6"] = "1) При подключении зарплатного проекта"
    table = Table(displayName="SalaryFees", ref="A2:B4")
    sheet.add_table(table)
    hidden = workbook.create_sheet("Скрытые условия")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Скрытый текст"
    workbook.save(path)
    return path


def _blocks_by_text(result):
    return {block.text: block for block in result.blocks}


def test_extracts_cells_tables_merged_headers_footnotes_hidden_and_percent(tmp_path: Path) -> None:
    path = _sample_xlsx(tmp_path / "sample.xlsx")
    result = XlsxParser().parse(_request(path), str(path))

    assert result.status == "ok"
    assert result.extraction_sha256
    by_text = _blocks_by_text(result)
    percent = by_text["0.3%"]
    assert percent.locator.kind == "xlsx"
    assert percent.locator.sheet == "Тарифы"
    assert percent.locator.cell_range == "B3"
    assert percent.table_id == "Тарифы:SalaryFees"
    assert {item["value"] for item in percent.relations if item["type"] == "column_header"} == {
        "Комиссия"
    }
    assert any(
        item.get("value") == "Зарплатный проект"
        for item in percent.relations
        if item["type"] == "merged_header"
    )
    assert "0,3%" in by_text
    assert any(
        item.get("range") == "A2:B4" for item in percent.relations if item["type"] == "table_cell"
    )

    merged = by_text["Зарплатный проект"]
    assert merged.locator.cell_range == "A1:B1"
    assert any(item["type"] == "merged_header" for item in merged.relations)
    assert "Скрытый текст" in by_text
    assert "Индивидуально" in by_text
    assert any(item.code == "XLSX_HIDDEN_SHEET" for item in result.diagnostics)
    assert any(item.code == "XLSX_MIXED_COLUMN_TYPES" for item in result.diagnostics)
    footnote = by_text["1) При подключении зарплатного проекта"]
    assert any(item["type"] == "footnote" for item in footnote.relations)
    assert any(item.code == "XLSX_FOOTNOTE_RELATION_UNVERIFIED" for item in result.diagnostics)
    formula = by_text["=1+1"]
    assert formula.locator.formula == "=1+1"
    assert formula.locator.cached_value is None
    assert any(item["type"] == "cached_value_missing" for item in formula.relations)
    assert any(item.code == "XLSX_FORMULA_CACHE_MISSING" for item in result.diagnostics)


def test_extraction_hash_and_block_ids_are_stable_and_original_hash_is_checked(
    tmp_path: Path,
) -> None:
    path = _sample_xlsx(tmp_path / "stable.xlsx")
    parser = XlsxParser()
    request = _request(path)
    first = parser.parse(request, str(path))
    second = parser.parse(request, str(path))
    assert first.extraction_sha256 == second.extraction_sha256
    assert [block.block_id for block in first.blocks] == [block.block_id for block in second.blocks]

    mismatch = request.model_copy(update={"raw_sha256": "0" * 64})
    rejected = parser.parse(mismatch, str(path))
    assert rejected.status == "failed"
    assert rejected.diagnostics[0].code == "XLSX_ORIGINAL_HASH_MISMATCH"


def test_formula_cache_is_preserved_without_evaluation(tmp_path: Path) -> None:
    source_path = _sample_xlsx(tmp_path / "source.xlsx")
    cached_path = tmp_path / "cached.xlsx"
    with (
        zipfile.ZipFile(source_path) as source,
        zipfile.ZipFile(cached_path, "w", compression=zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content, replacements = re.subn(
                    rb'(<c[^>]*r="B5"[^>]*><f>1\+1</f>)<v\s*/></c>',
                    rb"\g<1><v>2</v></c>",
                    content,
                )
                assert replacements == 1
            target.writestr(item, content)

    result = XlsxParser().parse(_request(cached_path), str(cached_path))
    formula = _blocks_by_text(result)["=1+1 [cached: 2]"]
    assert formula.locator.formula == "=1+1"
    assert formula.locator.cached_value == "2"
    assert not any(
        item.code == "XLSX_FORMULA_CACHE_MISSING" and item.block_id == formula.block_id
        for item in result.diagnostics
    )


def test_rejects_corrupt_archive_external_relationships_and_zip_bomb(tmp_path: Path) -> None:
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_bytes(b"PK\x03\x04not a workbook")
    assert XlsxParser().parse(_request(corrupt), str(corrupt)).status == "failed"

    valid = _sample_xlsx(tmp_path / "external.xlsx")
    rewritten = tmp_path / "external-rewritten.xlsx"
    with zipfile.ZipFile(valid) as source, zipfile.ZipFile(rewritten, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename.endswith(".rels") and b"TargetMode" not in content:
                content = content.replace(
                    b"</Relationships>",
                    (
                        b'<Relationship TargetMode="External" '
                        b'Target="https://example.invalid" Id="rId999" '
                        b'Type="external"/></Relationships>'
                    ),
                )
            target.writestr(item, content)
    external_result = XlsxParser().parse(_request(rewritten), str(rewritten))
    assert external_result.status == "failed"
    assert external_result.diagnostics[0].code == "XLSX_FORMAT_REJECTED"

    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x" * 2_000_000)
        archive.writestr("xl/workbook.xml", "x" * 2_000_000)
    bomb_result = XlsxParser().parse(_request(bomb), str(bomb))
    assert bomb_result.status == "failed"
    assert bomb_result.diagnostics[0].code == "XLSX_FORMAT_REJECTED"


def test_enforces_file_cell_and_dimension_limits(tmp_path: Path) -> None:
    path = _sample_xlsx(tmp_path / "limits.xlsx")
    request = _request(path)
    assert (
        XlsxParser(max_file_bytes=1).parse(request, str(path)).diagnostics[0].code
        == "XLSX_FILE_SIZE_LIMIT"
    )
    assert (
        XlsxParser(max_cells=1).parse(request, str(path)).diagnostics[0].code == "XLSX_RECTANGLE_LIMIT"
    )

    workbook = Workbook()
    workbook.active["ZZ1"] = "far column"
    wide = tmp_path / "wide.xlsx"
    workbook.save(wide)
    result = XlsxParser(max_columns=10).parse(_request(wide), str(wide))
    assert result.status == "failed"
    assert result.diagnostics[0].code == "XLSX_DIMENSION_LIMIT"


def test_sparse_worksheet_cannot_materialize_an_unbounded_rectangle(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "first"
    workbook.active["SF100000"] = "sparse last cell"
    path = tmp_path / "sparse.xlsx"
    workbook.save(path)
    result = XlsxParser().parse(_request(path), str(path))
    assert result.status == "failed"
    assert result.blocks == ()
    assert result.diagnostics[0].code == "XLSX_RECTANGLE_LIMIT"


def test_percent_display_retains_underlying_numeric_precision(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = 0.00333
    workbook.active["A1"].number_format = "0.0%"
    path = tmp_path / "precision.xlsx"
    workbook.save(path)
    result = XlsxParser().parse(_request(path), str(path))
    assert result.status == "ok"
    block = result.blocks[0]
    assert block.text == "0.3%"
    assert any(r["type"] == "numeric_cell" and r["raw_value"] == "0.00333"
               for r in block.relations)
