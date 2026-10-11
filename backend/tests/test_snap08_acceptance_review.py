"""Independent SNAP-08 acceptance regressions for the final review findings."""

import json
from io import BytesIO
from urllib.parse import quote

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from sqlalchemy import event, func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from test_extraction_pipeline import (
    _CONTEXT,
    HTML_VERSION,
    _html_bytes,
    _mapping_for_block,
    _seed_document,
)
from test_extraction_pipeline import pipeline_db as pipeline_db

from app.corpus.models import DocumentChunk, ExtractionRun
from app.corpus.service import CorpusError
from app.parsing.service import ExtractionPipelineError
from app.registry.models import SourcePolicy
from app.tariffs.models import EvidenceRef, TariffCandidate


@pytest.mark.parametrize("text,mapped", [
    ("Обслуживание счета от 100 RUB в месяц.", False),
    ("Обслуживание счета в месяц — не бесплатно.", False),
    ("100 RUB или 200 RUB", True),
    ("Обслуживание счета -100 RUB в месяц.", False),
    ("Обслуживание счета 100–200 RUB в месяц.", False),
    ("Обслуживание счета <100 RUB в месяц.", True),
])
def test_ambiguous_or_negated_value_is_not_a_positive_exact_fee(pipeline_db, text, mapped):
    ids, manifest, parsed, pipeline = _seed_document(pipeline_db, "html", _html_bytes(text))
    result = pipeline.run(
        manifest, ids["version"], parser_version=HTML_VERSION, context=_CONTEXT,
        operator_mapping=_mapping_for_block(parsed, kind="html") if mapped else None,
    )
    with pipeline_db[0]() as session:
        rows = session.scalars(select(TariffCandidate).where(
            TariffCandidate.extraction_run_id == result.extraction_id
        )).all()
        assert not any(row.payload["value"]["kind"] == "money" for row in rows), [
            row.payload for row in rows
        ]


def test_merged_header_that_selects_metric_is_attached_as_evidence(pipeline_db):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Обслуживание счета в месяц"
    sheet.append(["Услуга", "Стоимость"])
    sheet.append(["Пакет", "100 RUB"])
    sheet.add_table(Table(displayName="RatesTable", ref="A2:B3"))
    output = BytesIO()
    workbook.save(output)
    ids, manifest, parsed, pipeline = _seed_document(pipeline_db, "xlsx", output.getvalue())
    result = pipeline.run(
        manifest, ids["version"], parser_version=parsed.parser_version, context=_CONTEXT,
    )
    assert result.candidate_count == 1
    with pipeline_db[0]() as session:
        refs = session.scalars(select(EvidenceRef).where(
            EvidenceRef.extraction_run_id == result.extraction_id
        )).all()
        assert any(ref.locator.get("cell_range") == "A1:B1" for ref in refs), [
            (ref.quote, ref.locator, ref.support_role) for ref in refs
        ]


def test_invalid_mapping_can_be_corrected_without_faking_parser_version(pipeline_db):
    ids, manifest, parsed, pipeline = _seed_document(pipeline_db, "html", _html_bytes("100 RUB"))
    invalid = json.dumps({"version": 1, "rules": [{
        "metric_code": "account_monthly_fee", "value_block_id": "missing",
    }]})
    initial = pipeline.run(
        manifest, ids["version"], parser_version=HTML_VERSION,
        operator_mapping=invalid, context=_CONTEXT,
    )
    assert initial.candidate_count == 0
    assert initial.status == "failed" and initial.extraction_id is None
    corrected = pipeline.run(
        manifest, ids["version"], parser_version=HTML_VERSION,
        operator_mapping=_mapping_for_block(parsed, kind="html"), context=_CONTEXT,
    )
    assert corrected.candidate_count == 1


@pytest.mark.parametrize("missing_context", [False, True])
def test_incomplete_operator_input_can_be_repaired(pipeline_db, missing_context):
    text = "Обслуживание счета — 100 RUB в месяц." if missing_context else "100 RUB"
    ids, manifest, parsed, pipeline = _seed_document(pipeline_db, "html", _html_bytes(text))
    initial = pipeline.run(
        manifest, ids["version"], parser_version=HTML_VERSION,
        context=None if missing_context else _CONTEXT,
    )
    assert initial.status == "failed" and initial.extraction_id is None
    corrected = pipeline.run(
        manifest, ids["version"], parser_version=HTML_VERSION,
        operator_mapping=_mapping_for_block(parsed, kind="html"), context=_CONTEXT,
    )
    assert corrected.status == "ok" and corrected.candidate_count == 1
    assert pipeline.run(
        manifest, ids["version"], parser_version=HTML_VERSION,
        operator_mapping=_mapping_for_block(parsed, kind="html"), context=_CONTEXT,
    ).idempotent


def test_database_failure_after_candidate_flush_rolls_back_entire_chain(pipeline_db):
    ids, manifest, _, pipeline = _seed_document(pipeline_db, "html", _html_bytes())

    def fail_on_evidence(session, flush_context, instances):
        if any(isinstance(row, EvidenceRef) for row in session.new):
            raise SQLAlchemyError("controlled acceptance fault")

    event.listen(Session, "before_flush", fail_on_evidence)
    try:
        with pytest.raises(ExtractionPipelineError, match="EXTRACTION_PERSISTENCE_FAILED"):
            pipeline.run(manifest, ids["version"], parser_version=HTML_VERSION, context=_CONTEXT)
    finally:
        event.remove(Session, "before_flush", fail_on_evidence)
    with pipeline_db[0]() as session:
        for model in (ExtractionRun, DocumentChunk, TariffCandidate, EvidenceRef):
            assert session.scalar(select(func.count()).select_from(model)) == 0


def test_permission_revoked_during_parse_is_rechecked_before_write(pipeline_db):
    ids, manifest, _, pipeline = _seed_document(pipeline_db, "html", _html_bytes())
    original_runner = pipeline.parse_runner

    def revoke_after_parse(*args, **kwargs):
        parsed = original_runner(*args, **kwargs)
        with pipeline_db[0].begin() as session:
            session.get(SourcePolicy, ids["policy"]).collection_decision = "denied"
        return parsed

    pipeline.parse_runner = revoke_after_parse
    with pytest.raises(CorpusError, match="PERMISSION_REVOKED"):
        pipeline.run(manifest, ids["version"], parser_version=HTML_VERSION, context=_CONTEXT)
    with pipeline_db[0]() as session:
        assert session.scalar(select(func.count()).select_from(ExtractionRun)) == 0


@pytest.mark.parametrize("marker_on_label", [False, True])
def test_table_footnote_on_label_or_value_is_condition_evidence(pipeline_db, marker_on_label):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Услуга", "Стоимость"])
    sheet.append([
        "Обслуживание счета в месяц" + ("*" if marker_on_label else ""),
        "100 RUB" + ("" if marker_on_label else "*"),
    ])
    sheet["A4"] = "* При поддержании остатка."
    sheet.add_table(Table(displayName="RatesTable", ref="A1:B2"))
    output = BytesIO()
    workbook.save(output)
    ids, manifest, parsed, pipeline = _seed_document(pipeline_db, "xlsx", output.getvalue())
    result = pipeline.run(
        manifest, ids["version"], parser_version=parsed.parser_version, context=_CONTEXT,
    )
    assert result.candidate_count == 1
    with pipeline_db[0]() as session:
        refs = session.scalars(select(EvidenceRef).where(
            EvidenceRef.extraction_run_id == result.extraction_id
        )).all()
        assert any(ref.locator.get("cell_range") == "A4" and ref.support_role == "condition"
                   for ref in refs), [(ref.quote, ref.support_role) for ref in refs]


@pytest.mark.parametrize("target_id", ["note", "условие"])
def test_html_footnote_fragment_is_condition_evidence(pipeline_db, target_id):
    body = (
        '<html><body><table><tr><td>Обслуживание счета в месяц</td>'
        f'<td>100 RUB<a href="#{quote(target_id)}">*</a></td></tr></table>'
        f'<p id="{target_id}" role="doc-footnote">* При поддержании остатка.</p></body></html>'
    ).encode()
    ids, manifest, _, pipeline = _seed_document(pipeline_db, "html", body)
    result = pipeline.run(
        manifest, ids["version"], parser_version=HTML_VERSION, context=_CONTEXT,
    )
    assert result.candidate_count == 1
    with pipeline_db[0]() as session:
        refs = session.scalars(select(EvidenceRef).where(
            EvidenceRef.extraction_run_id == result.extraction_id
        )).all()
        assert any(ref.support_role == "condition" and "остатка" in ref.quote
                   for ref in refs), [(ref.quote, ref.support_role) for ref in refs]


@pytest.mark.parametrize("target", ["missing", "duplicate", "duplicate_unextracted"])
def test_unresolved_html_footnote_does_not_create_numeric_candidate(pipeline_db, target):
    notes = "" if target == "missing" else (
        '<p id="note" role="doc-footnote">* При остатке.</p>'
        + ('<div id="note">Unknown target</div>' if target == "duplicate_unextracted" else
           '<p id="note" role="doc-footnote">* Только по договору.</p>')
    )
    body = (
        '<html><body><table><tr><td>Обслуживание счета в месяц</td>'
        '<td>100 RUB<a href="#note">*</a></td></tr></table>' + notes + '</body></html>'
    ).encode()
    ids, manifest, _, pipeline = _seed_document(pipeline_db, "html", body)
    result = pipeline.run(
        manifest, ids["version"], parser_version=HTML_VERSION, context=_CONTEXT,
    )
    assert result.status == "failed" and result.candidate_count == 0
    assert any(item.code == "FOOTNOTE_REFERENCE_UNRESOLVED" for item in result.diagnostics)
