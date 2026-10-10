"""Deterministic, bounded extraction from verified local XLSX originals."""

from __future__ import annotations

import hashlib
import json
import os
import re
import warnings
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from uuid import NAMESPACE_URL, uuid5
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import coordinate_from_string, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from app.contracts.parsing import (
    ParsedBlock,
    ParseDiagnostic,
    Parser,
    ParseRequest,
    ParseResult,
    XlsxLocator,
)
from app.ingestion.errors import IngestionError
from app.ingestion.formats import check_format

PARSER_VERSION = "openpyxl/2"
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_SHEETS = 100
MAX_CELLS = 200_000
MAX_ROWS = 100_000
MAX_COLUMNS = 500
MAX_TEXT_CHARS = 10_000_000
MAX_PACKAGE_TEXT_CHARS = 20_000_000
_FOOTNOTE = re.compile(r"^\s*(?:(\d{1,3})[.)]|\[(\d{1,3})\]|([*†‡]))\s+")
_MARKER_REFERENCE = re.compile(r"(?:\[(\d{1,3})\]|(?<!\d)(\d{1,3})[†‡]|\*(?!\w))")
_XML_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class XlsxParser(Parser):
    """Parse a previously verified local XLSX without evaluating formulas or links."""

    def __init__(
        self,
        *,
        max_file_bytes: int = MAX_FILE_BYTES,
        max_sheets: int = MAX_SHEETS,
        max_cells: int = MAX_CELLS,
        max_rows: int = MAX_ROWS,
        max_columns: int = MAX_COLUMNS,
        max_text_chars: int = MAX_TEXT_CHARS,
    ) -> None:
        self.max_file_bytes = max_file_bytes
        self.max_sheets = max_sheets
        self.max_cells = max_cells
        self.max_rows = max_rows
        self.max_columns = max_columns
        self.max_text_chars = max_text_chars

    def parse(self, request: ParseRequest, original_path: str) -> ParseResult:
        if request.kind != "xlsx":
            return self._result(request, "unsupported", "XLSX_KIND_REQUIRED")
        path = Path(original_path)
        try:
            with path.open("rb") as original:
                size = os.fstat(original.fileno()).st_size
                if size > self.max_file_bytes:
                    return self._result(request, "failed", "XLSX_FILE_SIZE_LIMIT")
                digest = _sha256_stream(original)
                if digest != request.raw_sha256:
                    return self._result(request, "failed", "XLSX_ORIGINAL_HASH_MISMATCH")
            # Keep ingestion's established container and policy checks before the parser.
            check_format(path, "xlsx")
            package_stats = self._preflight(path)
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                formulas_wb = load_workbook(
                    path, read_only=False, data_only=False, keep_links=False
                )
                cached_wb = load_workbook(path, read_only=False, data_only=True, keep_links=False)
            try:
                if len(formulas_wb.worksheets) > self.max_sheets:
                    return self._result(request, "failed", "XLSX_SHEET_LIMIT")
                return self._extract(request, formulas_wb, cached_wb, package_stats)
            finally:
                formulas_wb.close()
                cached_wb.close()
        except _ResourceLimit as exc:
            return self._result(request, "failed", exc.code)
        except IngestionError:
            return self._result(request, "failed", "XLSX_FORMAT_REJECTED")
        except (
            InvalidFileException,
            zipfile.BadZipFile,
            ET.ParseError,
            OSError,
            ValueError,
        ) as exc:
            return self._result(request, "failed", "XLSX_CORRUPT_OR_UNREADABLE", type(exc).__name__)
        except Exception as exc:  # openpyxl can raise parser-specific exceptions for malformed XML
            return self._result(request, "failed", "XLSX_PARSE_ERROR", type(exc).__name__)

    def _preflight(self, path: Path) -> dict[str, int]:
        """Bound XML work and worksheet dimensions before openpyxl materializes cells."""
        cell_count = 0
        text_chars = 0
        sheet_xml_count = 0
        merged_cell_count = 0
        rectangle_cells = 0
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
                raise _ResourceLimit("XLSX_PACKAGE_SIZE_LIMIT")
            for entry in entries:
                is_sheet = entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(
                    ".xml"
                )
                is_shared_strings = entry.filename == "xl/sharedStrings.xml"
                if is_sheet:
                    sheet_xml_count += 1
                    if sheet_xml_count > self.max_sheets:
                        raise _ResourceLimit("XLSX_SHEET_LIMIT")
                if not is_sheet and not is_shared_strings:
                    continue
                with archive.open(entry) as stream:
                    max_row_seen = max_col_seen = 0
                    for _event, element in ET.iterparse(stream, events=("end",)):
                        if is_sheet and element.tag == f"{_XML_NS}dimension":
                            bounds = _required_bounds(element.attrib.get("ref", "A1"))
                            if bounds[2] > self.max_columns or bounds[3] > self.max_rows:
                                raise _ResourceLimit("XLSX_DIMENSION_LIMIT")
                            if bounds[2] * bounds[3] > self.max_cells:
                                raise _ResourceLimit("XLSX_RECTANGLE_LIMIT")
                        if is_sheet and element.tag == f"{_XML_NS}c":
                            cell_count += 1
                            if cell_count > self.max_cells:
                                raise _ResourceLimit("XLSX_CELL_LIMIT")
                            ref = element.attrib.get("r", "")
                            if ref:
                                try:
                                    col_letters, row = coordinate_from_string(ref)
                                    col = _column_number(col_letters)
                                except ValueError as exc:
                                    raise _ResourceLimit("XLSX_INVALID_CELL_REFERENCE") from exc
                                if row > self.max_rows or col > self.max_columns:
                                    raise _ResourceLimit("XLSX_DIMENSION_LIMIT")
                                max_row_seen = max(max_row_seen, row)
                                max_col_seen = max(max_col_seen, col)
                        elif is_sheet and element.tag == f"{_XML_NS}mergeCell":
                            ref = element.attrib.get("ref", "")
                            min_col, min_row, max_col, max_row = _required_bounds(ref)
                            if max_row > self.max_rows or max_col > self.max_columns:
                                raise _ResourceLimit("XLSX_DIMENSION_LIMIT")
                            merged_cell_count += (max_row - min_row + 1) * (max_col - min_col + 1)
                            if merged_cell_count > self.max_cells:
                                raise _ResourceLimit("XLSX_MERGED_RANGE_LIMIT")
                        if element.tag in (f"{_XML_NS}t", f"{_XML_NS}v", f"{_XML_NS}f"):
                            text_chars += len(element.text or "")
                            if text_chars > MAX_PACKAGE_TEXT_CHARS:
                                raise _ResourceLimit("XLSX_PACKAGE_TEXT_LIMIT")
                        element.clear()
                    rectangle_cells += max_row_seen * max_col_seen
                    if rectangle_cells > self.max_cells:
                        raise _ResourceLimit("XLSX_RECTANGLE_LIMIT")
            if sheet_xml_count == 0:
                raise _ResourceLimit("XLSX_NO_WORKSHEETS")
        return {"cells": cell_count, "text_chars": text_chars, "sheets": sheet_xml_count}

    def _extract(
        self,
        request: ParseRequest,
        formulas_wb: Any,
        cached_wb: Any,
        package_stats: dict[str, int],
    ) -> ParseResult:
        blocks: list[ParsedBlock] = []
        diagnostics: list[ParseDiagnostic] = []
        text_count = 0
        sheet_map = {sheet.title: sheet for sheet in cached_wb.worksheets}
        for formulas_sheet in formulas_wb.worksheets:
            if formulas_sheet.max_row * formulas_sheet.max_column > self.max_cells:
                raise _ResourceLimit("XLSX_RECTANGLE_LIMIT")
            cached_sheet = sheet_map[formulas_sheet.title]
            hidden = formulas_sheet.sheet_state != "visible"
            if hidden:
                diagnostics.append(
                    ParseDiagnostic(
                        code="XLSX_HIDDEN_SHEET",
                        message=(
                            f"Sheet {formulas_sheet.title!r} is {formulas_sheet.sheet_state}; "
                            "content retained."
                        ),
                    )
                )
            merged = list(formulas_sheet.merged_cells.ranges)
            merged_anchors: dict[str, str] = {}
            merged_header_ranges: list[tuple[int, int, int, int, str]] = []
            for merged_range in merged:
                min_col, min_row, max_col, max_row = _required_bounds(str(merged_range))
                anchor = formulas_sheet.cell(min_row, min_col)
                value = _cell_text(anchor)
                if value:
                    merged_anchors[anchor.coordinate] = str(merged_range)
                    if max_col > min_col or max_row > min_row:
                        merged_header_ranges.append((min_row, min_col, max_row, max_col, value))

            tables = list(formulas_sheet.tables.values())
            table_info: list[tuple[str, int, int, int, int, list[str], str]] = []
            for table in tables:
                min_col, min_row, max_col, max_row = _required_bounds(table.ref)
                headers = (
                    [
                        _cell_text(formulas_sheet.cell(min_row, col)) or ""
                        for col in range(min_col, max_col + 1)
                    ]
                    if table.headerRowCount
                    else []
                )
                table_info.append(
                    (table.displayName, min_row, min_col, max_row, max_col, headers, table.ref)
                )

            by_column: dict[int, set[str]] = {}
            sheet_blocks: list[ParsedBlock] = []
            current_heading: str | None = None
            for row in formulas_sheet.iter_rows():
                row_values = [
                    cell for cell in row if isinstance(cell, Cell) and cell.value is not None
                ]
                if len(row_values) == 1 and isinstance(row_values[0].value, str):
                    candidate_heading = row_values[0].value.strip()
                    if candidate_heading and not candidate_heading.startswith("="):
                        current_heading = candidate_heading
                for cell in row:
                    if not isinstance(cell, Cell):
                        continue
                    if cell.coordinate not in merged_anchors and _is_merged_placeholder(
                        cell, merged
                    ):
                        continue
                    formula = (
                        cell.value
                        if isinstance(cell.value, str) and cell.value.startswith("=")
                        else None
                    )
                    cached_cell = cached_sheet[cell.coordinate]
                    display = _cell_text(cell, formula_value=cached_cell.value if formula else None)
                    if not display:
                        continue
                    text_count += len(display)
                    if text_count > self.max_text_chars:
                        raise _ResourceLimit("XLSX_TEXT_LIMIT")
                    cached_value = (
                        _scalar_text(cached_cell.value)
                        if formula and cached_cell.value is not None
                        else None
                    )
                    locator_range = merged_anchors.get(cell.coordinate, cell.coordinate)
                    table = next(
                        (
                            item
                            for item in table_info
                            if item[1] <= cell.row <= item[3] and item[2] <= cell.column <= item[4]
                        ),
                        None,
                    )
                    relations: list[dict[str, str]] = []
                    if isinstance(cell.value, (int, float, Decimal)) and not isinstance(
                        cell.value, bool
                    ):
                        relations.append({
                            "type": "numeric_cell",
                            "raw_value": _scalar_text(cell.value),
                            "number_format": cell.number_format,
                        })
                    if hidden:
                        relations.append(
                            {"type": "hidden_sheet", "state": formulas_sheet.sheet_state}
                        )
                    if locator_range != cell.coordinate:
                        relations.append({"type": "merged_cell_range", "range": locator_range})
                    for start_row, start_col, end_row, end_col, header_text in merged_header_ranges:
                        if start_col <= cell.column <= end_col and (
                            cell.row > end_row or start_row <= cell.row <= end_row
                        ):
                            relations.append(
                                {
                                    "type": "merged_header",
                                    "value": header_text,
                                    "review_required": "true",
                                }
                            )
                    table_id = None
                    if table:
                        table_id = f"{formulas_sheet.title}:{table[0]}"
                        relations.append(
                            {"type": "table_cell", "table": table[0], "range": table[6]}
                        )
                        headers = table[5]
                        header_index = cell.column - table[2]
                        if (
                            headers
                            and cell.row > table[1]
                            and header_index < len(headers)
                            and headers[header_index]
                        ):
                            relations.append(
                                {
                                    "type": "column_header",
                                    "value": headers[header_index],
                                    "review_required": "true",
                                }
                            )
                    if formula:
                        relations.append({"type": "formula_not_evaluated"})
                        if cached_value is None:
                            relations.append(
                                {"type": "cached_value_missing", "review_required": "true"}
                            )
                            diagnostics.append(
                                ParseDiagnostic(
                                    code="XLSX_FORMULA_CACHE_MISSING",
                                    message=(
                                        f"Formula at {formulas_sheet.title}!{cell.coordinate} "
                                        "has no cached result; value remains unknown."
                                    ),
                                )
                            )
                    if _FOOTNOTE.match(display):
                        marker = _FOOTNOTE.match(display)
                        assert marker is not None
                        marker_value = next((value for value in marker.groups() if value), "*")
                        relations.append(
                            {"type": "footnote", "marker": marker_value, "review_required": "true"}
                        )
                    block = _make_block(
                        request,
                        formulas_sheet.title,
                        locator_range,
                        display,
                        formula=formula,
                        cached_value=cached_value,
                        heading=current_heading,
                        table_id=table_id,
                        relations=relations,
                        ordinal=len(blocks),
                    )
                    blocks.append(block)
                    sheet_blocks.append(block)
                    if not formula and not isinstance(cell.value, bool):
                        data_row = table is None or cell.row > table[1]
                        if data_row:
                            by_column.setdefault(cell.column, set()).add(_value_type(cell.value))

            for col, types in by_column.items():
                if len(types) > 1:
                    diagnostics.append(
                        ParseDiagnostic(
                            code="XLSX_MIXED_COLUMN_TYPES",
                            message=(
                                f"Sheet {formulas_sheet.title!r}, column {col} contains mixed "
                                f"value types: {', '.join(sorted(types))}."
                            ),
                        )
                    )
            diagnostics.extend(_link_footnotes(sheet_blocks))

        if not blocks:
            return self._result(request, "failed", "XLSX_NO_EXTRACTABLE_CELLS")
        canonical = json.dumps(
            [block.model_dump(mode="json") for block in blocks],
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        if package_stats["cells"] > self.max_cells:
            raise _ResourceLimit("XLSX_CELL_LIMIT")
        return ParseResult(
            status="ok",
            extraction_sha256=hashlib.sha256(canonical).hexdigest(),
            parser_version=request.parser_version,
            blocks=tuple(blocks),
            diagnostics=tuple(diagnostics),
        )

    @staticmethod
    def _result(request: ParseRequest, status: str, code: str, message: str = "") -> ParseResult:
        return ParseResult(
            status=status,  # type: ignore[arg-type]
            parser_version=request.parser_version,
            diagnostics=(ParseDiagnostic(code=code, message=message or code),),
        )


class _ResourceLimit(Exception):
    def __init__(self, code: str) -> None:
        self.code = code


def _make_block(
    request: ParseRequest,
    sheet: str,
    cell_range: str,
    text: str,
    *,
    formula: str | None,
    cached_value: str | None,
    heading: str | None,
    table_id: str | None,
    relations: list[dict[str, str]],
    ordinal: int,
) -> ParsedBlock:
    identity = (
        f"{request.raw_sha256}:{request.parser_version}:{sheet}:{cell_range}:{ordinal}:{text}"
    )
    return ParsedBlock(
        block_id=uuid5(NAMESPACE_URL, identity),
        text=text,
        locator=XlsxLocator(
            kind="xlsx",
            sheet=sheet,
            cell_range=cell_range,
            formula=formula,
            cached_value=cached_value,
        ),
        heading=heading,
        table_id=table_id,
        relations=tuple(relations),
    )


def _cell_text(cell: Cell, *, formula_value: Any = None) -> str:
    if isinstance(cell.value, str) and cell.value.startswith("="):
        return (
            f"{cell.value} [cached: {_scalar_text(formula_value)}]"
            if formula_value is not None
            else cell.value
        )
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(value))
        except InvalidOperation:
            return str(value)
        if "%" in cell.number_format:
            number *= Decimal(100)
            decimals = _percent_decimals(cell.number_format)
            if decimals is not None:
                number = number.quantize(Decimal(1).scaleb(-decimals))
            return f"{_decimal_text(number)}%"
        return _decimal_text(number)
    if hasattr(value, "isoformat"):
        return str(value.isoformat())
    return str(value).replace("\x00", "").strip()


def _scalar_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float, Decimal)):
        try:
            return _decimal_text(Decimal(str(value)))
        except InvalidOperation:
            return str(value)
    if hasattr(value, "isoformat"):
        return str(value.isoformat())
    return str(value).replace("\x00", "").strip()


def _decimal_text(value: Decimal) -> str:
    rendered = format(value, "f")
    return rendered.rstrip("0").rstrip(".") if "." in rendered else rendered


def _percent_decimals(number_format: str) -> int | None:
    match = re.search(r"\.(0+)", number_format.split("%", 1)[0])
    return len(match.group(1)) if match else None


def _column_number(letters: str) -> int:
    result = 0
    for char in letters.upper():
        result = result * 26 + ord(char) - ord("A") + 1
    return result


def _required_bounds(cell_range: str) -> tuple[int, int, int, int]:
    bounds = range_boundaries(cell_range)
    if any(value is None for value in bounds):
        raise ValueError("XLSX_INVALID_RANGE")
    min_col, min_row, max_col, max_row = bounds
    assert min_col is not None
    assert min_row is not None
    assert max_col is not None
    assert max_row is not None
    return min_col, min_row, max_col, max_row


def _value_type(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float, Decimal)):
        return "number"
    if hasattr(value, "isoformat"):
        return "date_or_time"
    return "text"


def _is_merged_placeholder(cell: Cell, merged_ranges: list[Any]) -> bool:
    return any(
        cell.coordinate in merged_range and cell.coordinate != merged_range.start_cell.coordinate
        for merged_range in merged_ranges
    )


def _link_footnotes(blocks: list[ParsedBlock]) -> list[ParseDiagnostic]:
    notes: dict[str, ParsedBlock] = {}
    for block in blocks:
        match = _FOOTNOTE.match(block.text)
        if match:
            marker = next((value for value in match.groups() if value), "*")
            notes[marker] = block
    diagnostics: list[ParseDiagnostic] = []
    for block in blocks:
        if block in notes.values():
            continue
        refs: list[dict[str, str]] = []
        for match in _MARKER_REFERENCE.finditer(block.text):
            marker = match.group(1) or match.group(2) or "*"
            note = notes.get(marker)
            if note:
                refs.append(
                    {
                        "type": "footnote_reference",
                        "marker": marker,
                        "related_block_id": str(note.block_id),
                        "review_required": "true",
                    }
                )
                diagnostics.append(
                    ParseDiagnostic(
                        code="XLSX_FOOTNOTE_RELATION_UNVERIFIED",
                        message=(
                            f"Footnote relation for marker {marker} requires manual confirmation."
                        ),
                        block_id=block.block_id,
                    )
                )
        if refs:
            index = blocks.index(block)
            blocks[index] = block.model_copy(update={"relations": block.relations + tuple(refs)})
    for marker, note in notes.items():
        assert isinstance(note.locator, XlsxLocator)
        diagnostics.append(
            ParseDiagnostic(
                code="XLSX_FOOTNOTE_UNVERIFIED",
                message=(
                    f"Footnote {marker} at {note.locator.sheet}!{note.locator.cell_range} "
                    "requires manual confirmation."
                ),
                block_id=note.block_id,
            )
        )
    return diagnostics


def _sha256_stream(source: Any) -> str:
    digest = hashlib.sha256()
    for chunk in iter(lambda: source.read(1024 * 1024), b""):
        digest.update(chunk)
    return digest.hexdigest()
